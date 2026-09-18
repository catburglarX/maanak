"""DOCX rendering.

The editable format. Same frozen snapshot as the PDF, so the two documents always
agree; only the container differs.

Accessibility: headings use Word's built-in heading styles so that a screen reader and
the navigation pane both see real structure, tables get a header row marked as such,
and the document language is set so a reader pronounces it correctly.
"""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ...domain import labels as domain_labels

RENDERER_VERSION = "1"

#: Labels come from the domain, so the editable document, the PDF and the screen
#: cannot describe the same state in three different ways.
label_for = domain_labels.label_for

GREY = RGBColor(0x46, 0x46, 0x3F)
RED = RGBColor(0x9A, 0x1F, 0x1F)


def _set_language(document: DocumentObject, language: str = "en-IN") -> None:
    """Declare the document language for assistive technology."""
    styles_element = document.styles.element
    run_properties = styles_element.find(qn("w:docDefaults"))
    if run_properties is None:
        return
    for language_element in run_properties.iter(qn("w:lang")):
        language_element.set(qn("w:val"), language)


def _mark_header_row(table: Any) -> None:
    """Repeat the header row across pages and expose it as a header to readers."""
    header = table.rows[0]
    properties = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _key_value_table(document: DocumentObject, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value or "Not recorded"
        for paragraph in cells[0].paragraphs:
            for run in paragraph.runs:
                run.bold = True


def render_docx(snapshot: dict[str, Any], *, report: dict[str, Any]) -> bytes:
    """Render the editable report."""
    document = Document()
    _set_language(document)

    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = f"Maanak · {report['reference']} · {report['issuing_workspace']}"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT

    footer = section.footer.paragraphs[0]
    footer.text = f"Content hash {report['snapshot_sha256']}"
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT

    inspection = snapshot.get("inspection", {})
    product = snapshot.get("product", {})
    people = snapshot.get("people", {})
    rule_versions = snapshot.get("rule_versions", {})

    document.add_heading("Packaged commodity inspection report", level=0)
    subtitle = document.add_paragraph()
    run = subtitle.add_run(
        f"Issued by {report['issuing_workspace']} · Report {report['reference']} · "
        f"Revision {report['revision']} · Issued {report['issued_at']}"
    )
    run.font.size = Pt(9)
    run.font.color.rgb = GREY

    # --- Decision -------------------------------------------------------
    document.add_heading("Recorded decision", level=1)
    decision = str(inspection.get("decision") or "not recorded")
    decision_label = label_for(decision) if inspection.get("decision") else "Not recorded"

    paragraph = document.add_paragraph()
    decision_run = paragraph.add_run(decision_label)
    decision_run.bold = True
    decision_run.font.size = Pt(13)
    if decision == "violation_found":
        decision_run.font.color.rgb = RED

    document.add_paragraph(inspection.get("decision_note") or "No reason recorded.")
    decided_by = people.get(str(inspection.get("decided_by_id")), {})
    attribution = document.add_paragraph()
    run = attribution.add_run(
        f"Recorded by {decided_by.get('name', 'unknown')} "
        f"({decided_by.get('role', 'unknown role')}) on {inspection.get('decided_at')}"
    )
    run.font.size = Pt(8.5)

    # --- Inspection and product ----------------------------------------
    document.add_heading("Inspection", level=1)
    _key_value_table(
        document,
        [
            ("Inspection reference", str(inspection.get("reference"))),
            ("Inspection date", str(inspection.get("inspection_date"))),
            ("Source", label_for(str(inspection.get("source", "")))),
            ("Jurisdiction", str(inspection.get("jurisdiction_name"))),
            ("Premises", str(inspection.get("premises_name") or "")),
            ("Premises address", str(inspection.get("premises_address") or "")),
            ("Marketplace", str(inspection.get("marketplace_name") or "")),
            ("Listing URL", str(inspection.get("listing_url") or "")),
            ("Batch reference", str(inspection.get("batch_reference") or "")),
        ],
    )

    document.add_heading("Product", level=1)
    identifiers = ", ".join(
        f"{item['scheme'].upper()} {item['value']}" for item in product.get("identifiers", [])
    )
    quantity = product.get("declared_net_quantity")
    _key_value_table(
        document,
        [
            ("Brand", str(product.get("brand"))),
            ("Product", str(product.get("name"))),
            ("Common or generic name", str(product.get("common_generic_name") or "")),
            ("Commodity category", label_for(str(product.get("commodity_category") or ""))),
            ("Package type", label_for(str(product.get("package_type") or ""))),
            (
                "Declared net quantity",
                f"{quantity} {product.get('declared_net_quantity_unit') or ''}".strip()
                if quantity
                else "",
            ),
            ("Identifiers", identifiers),
            ("Imported", "Yes" if product.get("is_imported") else "No"),
            ("Country of origin", str(product.get("country_of_origin") or "")),
        ],
    )

    responsible = product.get("responsible_parties", [])
    if responsible:
        document.add_heading("Responsible parties", level=2)
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        headers = table.rows[0].cells
        headers[0].text = "Role"
        headers[1].text = "Name"
        headers[2].text = "Address"
        _mark_header_row(table)
        for item in responsible:
            cells = table.add_row().cells
            cells[0].text = label_for(item["party_role"])
            cells[1].text = item["legal_name"]
            cells[2].text = ", ".join(
                part
                for part in (
                    item.get("address_line"),
                    item.get("locality"),
                    item.get("state"),
                    item.get("pin_code"),
                    item.get("country"),
                )
                if part
            )

    # --- Evidence -------------------------------------------------------
    document.add_heading("Evidence and chain of custody", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    headers = table.rows[0].cells
    headers[0].text = "Package face"
    headers[1].text = "Received"
    headers[2].text = "Size"
    headers[3].text = "SHA-256 of the original file"
    _mark_header_row(table)
    for item in snapshot.get("evidence", []):
        cells = table.add_row().cells
        cells[0].text = label_for(item["face"])
        cells[1].text = str(item.get("server_received_at"))[:19]
        cells[2].text = f"{item['size_bytes'] // 1024} KB"
        cells[3].text = item["sha256"]

    not_captured = [
        item for item in snapshot.get("package_faces", []) if item["capture_state"] != "captured"
    ]
    if not_captured:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(
            "Package faces not captured: "
            + "; ".join(
                f"{label_for(item['face'])} ({label_for(item['capture_state']).lower()}"
                + (f", {item['reason']}" if item.get("reason") else "")
                + ")"
                for item in not_captured
            )
        )
        run.font.size = Pt(8.5)

    # --- Readings -------------------------------------------------------
    document.add_heading("Declarations read and reviewed", level=1)
    document.add_paragraph(
        "The machine reading and the officer's review are recorded separately. A "
        "correction never replaces what the machine read."
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    headers = table.rows[0].cells
    headers[0].text = "Declaration"
    headers[1].text = "Machine read"
    headers[2].text = "Officer review"
    headers[3].text = "Used in tests"
    _mark_header_row(table)
    for item in snapshot.get("readings", []):
        machine = item.get("machine_reading") or {}
        corrected = item.get("officer_corrected_value") or {}
        review = label_for(item["review_state"])
        if corrected:
            review += f" to {corrected.get('display') or corrected.get('value')}"
        if item.get("correction_reason"):
            review += f". Reason: {item['correction_reason']}"
        cells = table.add_row().cells
        cells[0].text = label_for(item["declaration_type"])
        cells[1].text = str(machine.get("display") or machine.get("value") or "not read")
        cells[2].text = review
        cells[3].text = "Yes" if item.get("used_for_legal_tests") else "No"

    # --- Findings -------------------------------------------------------
    document.add_heading("Findings", level=1)
    findings = snapshot.get("findings", [])
    if not findings:
        document.add_paragraph("No applicable rule produced a finding.")
    for index, finding in enumerate(findings, start=1):
        rule = rule_versions.get(finding["rule_version_id"], {})
        outcome = finding["effective_outcome"]
        document.add_heading(
            f"{index}. {rule.get('title', 'Rule')}: {label_for(outcome)}",
            level=2,
        )
        document.add_paragraph(finding["explanation"])
        rows = [
            ("Rule cited", str(rule.get("citation"))),
            (
                "Rule version",
                f"{rule.get('code')} version {rule.get('version')}, in force from "
                f"{rule.get('effective_from')}",
            ),
            (
                "Legal authority",
                "Confirmed by a qualified authority"
                if rule.get("legal_authority_confirmed")
                else "NOT confirmed by a statutory authority. Approved for use in this "
                "workspace only.",
            ),
        ]
        if finding.get("expected_value"):
            rows.append(("Expected", str(finding["expected_value"])))
        if finding.get("observed_value"):
            rows.append(("Observed", str(finding["observed_value"])))
        if finding.get("officer_outcome"):
            rows.append(
                (
                    "Officer override",
                    f"Engine returned {label_for(finding['engine_outcome'])}; "
                    f"officer recorded {label_for(finding['officer_outcome'])}. "
                    f"{finding.get('officer_note') or ''}",
                )
            )
        _key_value_table(document, rows)

        if finding.get("calculation"):
            document.add_paragraph("Calculation:")
            for step in finding["calculation"]:
                paragraph = document.add_paragraph(step, style="List Bullet")
                for run in paragraph.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(8.5)

    # --- Timeline -------------------------------------------------------
    document.add_heading("Timeline", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    headers = table.rows[0].cells
    headers[0].text = "When"
    headers[1].text = "Change"
    headers[2].text = "By"
    headers[3].text = "Reason"
    _mark_header_row(table)
    for item in snapshot.get("timeline", []):
        actor = people.get(str(item.get("actor_id")), {})
        cells = table.add_row().cells
        cells[0].text = str(item.get("occurred_at"))[:19]
        cells[1].text = (
            f"{label_for(item['from_state']) if item.get('from_state') else 'Opened'}"
            f" to {label_for(item['to_state'])}"
        )
        cells[2].text = actor.get("name") or item.get("actor_role") or "system"
        cells[3].text = item.get("reason") or "None given"

    # --- Verification ---------------------------------------------------
    document.add_heading("Verification", level=1)
    _key_value_table(
        document,
        [
            ("Report content hash", report["snapshot_sha256"]),
            ("Hash algorithm", report["snapshot_algorithm"]),
            ("Verification code", report["verification_code"]),
            ("Verify at", report.get("verification_url", "")),
            (
                "Digital signature",
                {
                    "none": "Not signed. This document carries a content hash only.",
                    "development": "Signed with a DEVELOPMENT key. Not valid for any "
                    "official purpose.",
                }.get(report["signature_status"], report["signature_status"]),
            ),
        ],
    )

    document.add_heading("Limits of this report", level=1)
    for limit in snapshot.get("limits", []):
        document.add_paragraph(limit, style="List Bullet")

    closing = document.add_paragraph()
    run = closing.add_run(
        "Maanak is not a government service and this report is not a government "
        "certification. It records an inspection carried out by the workspace named "
        "above."
    )
    run.font.size = Pt(8.5)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
